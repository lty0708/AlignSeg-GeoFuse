from pathlib import Path
import argparse
import csv
from dataclasses import asdict
import hashlib
import json
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent / 'src'))
import cv2
import nibabel as nib
import numpy as np
from openpyxl import load_workbook
from pitnet import io
from pitnet_m6_prepare import split_bilateral_carotids

def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')

def discover(data_root, metadata):
    wb = load_workbook(metadata, read_only=True, data_only=True)
    table = list(wb.active.iter_rows(values_only=True))
    header = list(table[0])
    pid_col, seq_col = (header.index('PID'), header.index('T1 sequence'))
    knosp_col = header.index('Knosp')
    output = []
    for values in table[1:]:
        if values[pid_col] is None:
            continue
        pid = int(values[pid_col])
        sequence = str(values[seq_col])
        files = {'image': data_root / f'{pid}_T1.nii.gz'}
        reasons = []
        for label in ('tumor', 'carotids'):
            candidates = [data_root / f'{pid}_{label}_f.nii.gz', data_root / f'{pid}_{label}.nii.gz']
            files[label] = next((p for p in candidates if p.is_file() and p.stat().st_size), next((p for p in candidates if p.is_file()), candidates[0]))
        for label, path in files.items():
            if not path.is_file():
                reasons.append(f'missing_{label}')
            elif path.stat().st_size == 0:
                reasons.append(f'zero_byte_{label}')
        if 'T2' in sequence.upper():
            reasons.append('metadata_T2_only')
        raw_knosp = str(values[knosp_col]).strip().upper()
        if raw_knosp not in {'0', '1', '2', '3', '3A', '3B', '4'}:
            reasons.append('missing_public_knosp')
        output.append({'pid': pid, 'sequence': sequence, 'eligible': not reasons, 'raw_knosp': raw_knosp, 'reasons': reasons, 'files': {k: str(v) for k, v in files.items()}})
    wb.close()
    return output

def prepare_case(row, out):
    files = {k: Path(v) for k, v in row['files'].items()}
    original, canonical, working, resampled = io.prepare_working_image(files['image'], (0.5,) * 3)
    tumor, tumor_mode, tumor_diff = io.align_mask_to_reference(files['tumor'], canonical, working)
    carotids, vessel_mode, vessel_diff = io.align_mask_to_reference(files['carotids'], canonical, working)
    right, left = split_bilateral_carotids(carotids)
    coords = np.argwhere(tumor | right | left)
    if not coords.size:
        raise ValueError('Empty reference foreground; cannot select planes')
    center = np.round(coords.mean(axis=0)).astype(int)
    axes = tuple((io.compute_axis_crop(n, 192, c) for n, c in zip(working.shape, center)))
    crop = io.crop_and_pad(np.asanyarray(working.dataobj), axes, (192,) * 3, np.float32)
    crop, low, high = io.robust_normalize(crop)
    cubes = {k: io.crop_and_pad(v, axes, (192,) * 3, np.uint8) for k, v in {'tumor': tumor, 'right_ICA': right, 'left_ICA': left}.items()}
    getters = {'m-1': lambda a: np.flip(a[:, 92, :].T, (0, 1)), 'm': lambda a: np.flip(a[:, 96, :].T, (0, 1)), 'm+1': lambda a: np.flip(a[:, 100, :].T, (0, 1)), 'b': lambda a: np.flip(a[96, :, :].T, 0), 'v': lambda a: np.flip(a[:, :, 96].T, 0)}
    out.mkdir(parents=True, exist_ok=False)
    masks = {}
    hashes = {}
    for view, getter in getters.items():
        image = np.round(np.clip(cv2.resize(getter(crop), (384, 384), interpolation=cv2.INTER_LINEAR), 0, 1) * 255).astype(np.uint8)
        ok, encoded = cv2.imencode('.png', image)
        assert ok
        encoded.tofile(str(out / f'{view}.png'))
        hashes[view] = hashlib.sha256(image.tobytes()).hexdigest()
        for name, cube in cubes.items():
            masks[f'{view}_{name}'] = (cv2.resize(getter(cube).astype(np.float32), (384, 384), interpolation=cv2.INTER_NEAREST) > 0.5).astype(np.uint8)
    np.savez_compressed(out / 'masks.npz', **masks)
    for name, mask in masks.items():
        hashes[name] = hashlib.sha256(mask.tobytes()).hexdigest()
    meta = {'pid': row['pid'], 'source_filenames': {k: p.name for k, p in files.items()}, 'original_axes': nib.aff2axcodes(original.affine), 'canonical_axes': nib.aff2axcodes(canonical.affine), 'original_affine': original.affine.tolist(), 'working_affine': working.affine.tolist(), 'original_spacing': list(map(float, original.header.get_zooms()[:3])), 'working_spacing_mm': [0.5] * 3, 'resampled': resampled, 'image_interpolation': 'linear', 'mask_interpolation': 'nearest', 'mask_alignment': {'tumor': tumor_mode, 'carotids': vessel_mode}, 'mask_affine_max_difference': {'tumor': tumor_diff, 'carotids': vessel_diff}, 'affine_policy_warning': 'Historical same-shape affine override retained; clinical correctness needs source audit when headers conflict.', 'working_center_xyz': center.tolist(), 'center_world_ras_mm': nib.affines.apply_affine(working.affine, center).tolist(), 'coronal_y_indices': [int(center[1] - 4), int(center[1]), int(center[1] + 4)], 'coronal_offset_mm': 2.0, 'crop_axes': [asdict(a) for a in axes], 'cube_size': 192, 'output_size': 384, 'normalization_percentiles': [0.5, 99.5], 'normalization_bounds': [low, high], 'scale_ruler_added': False, 'side_rule': 'RAS X larger is anatomical right; coronal display flips both axes after transpose', 'array_sha256': hashes}
    write_json(out / 'meta.json', meta)
    return meta

def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--data-root', type=Path, required=True)
    p.add_argument('--metadata', type=Path, required=True)
    p.add_argument('--output-dir', type=Path, required=True)
    p.add_argument('--audit-only', action='store_true')
    p.add_argument('--max-cases', type=int, default=None)
    args = p.parse_args()
    if args.output_dir.exists() and any(args.output_dir.iterdir()):
        raise FileExistsError('Use a new output directory; historical products are never overwritten')
    rows = discover(args.data_root, args.metadata)
    public_rows = [{**r, 'files': {k: Path(v).name for k, v in r['files'].items()}} for r in rows]
    summary = {'discovered': len(rows), 'eligible': sum((r['eligible'] for r in rows)), 'excluded': [r for r in public_rows if not r['eligible']], 'metadata_sha256': hashlib.sha256(args.metadata.read_bytes()).hexdigest()}
    write_json(args.output_dir / 'eligibility.json', summary)
    write_json(args.output_dir / 'inventory.json', public_rows)
    if not args.audit_only:
        selected = [r for r in rows if r['eligible']]
        if args.max_cases is not None:
            selected = selected[:args.max_cases]
        results = []
        for row in selected:
            try:
                meta = prepare_case(row, args.output_dir / 'cases' / f"{row['pid']:03d}")
                results.append({'pid': row['pid'], 'status': 'PASS', 'array_sha256': meta['array_sha256']})
            except Exception as e:
                results.append({'pid': row['pid'], 'status': 'FAIL', 'error': str(e)})
            print(json.dumps({k: v for k, v in results[-1].items() if k != 'array_sha256'}), flush=True)
        write_json(args.output_dir / 'preparation_results.json', results)
        if any((r['status'] != 'PASS' for r in results)):
            raise RuntimeError('Preparation failures retained in preparation_results.json')
    print(json.dumps(summary), flush=True)
if __name__ == '__main__':
    main()
